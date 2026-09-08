"""Build the editable manuscript with the bundled document runtime."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

HERE = Path(__file__).resolve().parent
DATA = json.loads((HERE / "document_data.json").read_text(encoding="utf-8"))
REFERENCES = json.loads((HERE / "references.json").read_text(encoding="utf-8"))


def element(tag, text=None, **attrs):
    item = OxmlElement(tag)
    if text is not None:
        item.text = text
    for key, value in attrs.items():
        item.set(qn(key), str(value))
    return item


def field(p, instruction):
    r = p.add_run()
    r._r.append(element("w:fldChar", **{"w:fldCharType": "begin"}))
    r._r.append(element("w:instrText", instruction, **{"xml:space": "preserve"}))
    r._r.append(element("w:fldChar", **{"w:fldCharType": "end"}))


def hyperlink(p, text, url):
    rel = p.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    h = element("w:hyperlink", **{"r:id": rel})
    r = element("w:r")
    pr = element("w:rPr")
    pr.append(element("w:color", **{"w:val": "000000"}))
    r.append(pr)
    r.append(element("w:t", text))
    h.append(r)
    p._p.append(h)


def add_text(p, text):
    # Keep URLs clickable; long artifact digests wrap without changing visible text.
    parts = re.split(r"(https?://\S+)", text)
    for part in parts:
        if part.startswith("http"):
            trail = "." if part.endswith(".") else ""
            url = part[:-1] if trail else part
            hyperlink(p, url, url)
            if trail:
                p.add_run(trail)
        else:
            for bit in re.split(r"([a-f0-9]{64})", part):
                if re.fullmatch("[a-f0-9]{64}", bit):
                    bit = "\u200b".join(bit[i : i + 8] for i in range(0, 64, 8))
                for token in re.split(r"(\bT_c\b|\bD_A\b)", bit):
                    if token in ["T_c", "D_A"]:
                        base, subscript = token.split("_")
                        p.add_run(base)
                        p.add_run(subscript).font.subscript = True
                    else:
                        p.add_run(token)


def set_section(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11.6929 if landscape else 8.2677)
    section.page_height = Inches(8.2677 if landscape else 11.6929)
    section.left_margin = section.right_margin = Inches(0.80 if landscape else 0.95)
    section.top_margin = section.bottom_margin = Inches(0.65 if landscape else 0.80)
    section.header_distance = section.footer_distance = Inches(0.30)


def make_table(doc, key):
    data = DATA["tables"][key]
    if key == "methods":
        for suffix, rows in [("a", data["rows"][:10]), ("b", data["rows"][10:])]:
            newkey = "methods_" + suffix
            DATA["tables"][newkey] = {
                **data,
                "rows": rows,
                "font": 10,
                "caption": data["caption"] + (" continued" if suffix == "b" else ""),
                "note": data["note"]
                if suffix == "b"
                else "Temporal pipelines. Spectral, wavelet and geometric pipelines continue on the following page.",
            }
            if suffix == "b":
                doc.add_page_break()
            make_table(doc, newkey)
        return
    p = doc.add_paragraph(data["caption"], "Caption")
    p.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(data["headers"]))
    table.autofit = False
    props = table._tbl.tblPr
    borders = element("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        borders.append(
            element("w:" + side, **{"w:val": "single", "w:sz": "4", "w:color": "C8C8C8"})
        )
    props.append(borders)
    margins = element("w:tblCellMar")
    for side in ["top", "bottom"]:
        margins.append(element("w:" + side, **{"w:w": "65", "w:type": "dxa"}))
    for side in ["left", "right"]:
        margins.append(element("w:" + side, **{"w:w": "75", "w:type": "dxa"}))
    props.append(margins)
    for column, width in zip(table.columns, data["widths"], strict=False):
        column.width = Inches(width)
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(element("w:tblHeader", **{"w:val": "true"}))
    allrows = [data["headers"]] + data["rows"]
    for i, values in enumerate(allrows):
        cells = header.cells if i == 0 else table.add_row().cells
        tr = cells[0]._tc.getparent()
        tr.get_or_add_trPr().append(element("w:cantSplit"))
        for j, (cell, text, width) in enumerate(zip(cells, values, data["widths"], strict=False)):
            cell.width = Inches(width)
            cell.vertical_alignment = 1
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = i == 0
            p.paragraph_format.widow_control = True
            if i == 0:
                cell._tc.get_or_add_tcPr().append(element("w:shd", **{"w:fill": "E7E7E7"}))
            elif i % 2 == 0:
                cell._tc.get_or_add_tcPr().append(element("w:shd", **{"w:fill": "F8F8F8"}))
            # Numeric intervals on two deliberate lines make narrow columns legible.
            if key in ["clean", "stress", "coverage"] and j >= (2 if key == "coverage" else 1):
                text = text.replace(" [", "\n[")
                text = re.sub(r"-0\.0+(?=[,\]\s]|$)", lambda m: m[0][1:], text)
            run = p.add_run(text)
            run.font.size = Pt(data.get("font", 9.5))
            run.bold = i == 0
    p = doc.add_paragraph(data["note"], "Table Note")
    p.paragraph_format.keep_with_next = False


def add_equation(doc, index):
    # Canonical MathML-to-OMML output avoids incomplete delimiter conversion.
    groups = json.loads((HERE / "equations.omml.json").read_text(encoding="utf-8"))
    for xml in groups[index]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_together = True
        p._p.append(etree.fromstring(xml.encode("utf-8")))


def main(output):
    doc = Document()
    doc.core_properties.title = "Benchmarking classical memory and scaling estimators under model mismatch and contamination"
    doc.core_properties.author = "Davian R. Chin"
    doc.core_properties.subject = "Synthetic confirmation benchmark of audited lrdbench pipelines"
    doc.core_properties.keywords = "long-range dependence; Hurst exponent; simulation; coverage"
    doc.core_properties.comments = "Rebuilt from audited confirmation results on 8 September 2026."
    styles = doc.styles
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"]:
        styles[name].font.name = "Times New Roman"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        rpr = styles[name].element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        for key in list(rfonts.attrib):
            if "Theme" in key:
                del rfonts.attrib[key]
        for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
            rfonts.set(qn("w:" + key), "Times New Roman")
    # Strip inherited title rules and theme borders from the starting template.
    for style in styles:
        for b in list(style.element.iter(qn("w:pBdr"))):
            b.getparent().remove(b)
    normal = styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True
    styles["Title"].font.size = Pt(20)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(14)
    for name, size in [("Heading 1", 13), ("Heading 2", 11.5)]:
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].paragraph_format.space_before = Pt(12)
        styles[name].paragraph_format.space_after = Pt(6)
        styles[name].paragraph_format.keep_with_next = True
    styles["Caption"].font.size = Pt(11)
    styles["Caption"].font.italic = False
    styles["Caption"].font.bold = True
    styles["Caption"].paragraph_format.space_after = Pt(7)
    note = styles.add_style("Table Note", 1)
    note.base_style = normal
    note.font.size = Pt(9)
    note.paragraph_format.space_before = Pt(6)
    note.paragraph_format.space_after = Pt(6)
    note.paragraph_format.line_spacing = 1.0
    set_section(doc.sections[0])
    head = doc.sections[0].header.paragraphs[0]
    head.text = "Memory and scaling estimator benchmark"
    head.runs[0].font.name = "Times New Roman"
    head.runs[0].font.size = Pt(9)
    foot = doc.sections[0].footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run("Draft · 8 September 2026     ")
    field(foot, "PAGE")
    for run in foot.runs:
        run.font.size = Pt(9)
    doc.settings.element.append(element("w:updateFields", **{"w:val": "true"}))
    equation_index = 0
    manuscript = (HERE / "manuscript.md").read_text(encoding="utf-8")
    citations = {int(n) for n in re.findall(r"\[(\d+)\]", manuscript)}
    assert citations == {r["id"] for r in REFERENCES}, citations
    for block in manuscript.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block in ["PAGEBREAK", "LANDSCAPE", "PORTRAIT"]:
            if block == "PAGEBREAK":
                doc.add_page_break()
            else:
                set_section(doc.add_section(WD_SECTION_START.NEW_PAGE), block == "LANDSCAPE")
        elif block.startswith("TABLE:"):
            make_table(doc, block.partition(":")[2].strip())
        elif block.startswith("FIGURE:"):
            fig = DATA["figures"][block.partition(":")[2].strip()]
            doc.add_paragraph(fig["caption"], "Caption")
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.add_run().add_picture(str(HERE / fig["path"]), width=Inches(9.65))
            doc.add_paragraph(fig["note"], "Table Note")
        elif block.startswith("EQUATION:"):
            add_equation(doc, equation_index)
            equation_index += 1
        elif block == "REFERENCES":
            for ref in REFERENCES:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(7)
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                p.add_run(f"[{ref['id']}] {ref['citation']} ")
                hyperlink(p, "https://doi.org/" + ref["doi"], "https://doi.org/" + ref["doi"])
        elif block.startswith("# "):
            doc.add_paragraph(block[2:], "Title")
        elif block.startswith("## "):
            if block.startswith("## 1 Introduction"):
                doc.add_page_break()
            doc.add_paragraph(block[3:], "Heading 1")
        elif block.startswith("### "):
            doc.add_paragraph(block[4:], "Heading 2")
        else:
            p = doc.add_paragraph()
            add_text(p, block)
    assert equation_index == 4
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    report = {
        "docx": str(output),
        "sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "tables": len(doc.tables),
        "pictures": len(doc.inline_shapes),
        "references": len(REFERENCES),
        "native_equation_groups": equation_index,
        "unresolved_tokens": bool(re.search(r"\{\{\w+\}\}", manuscript)),
    }
    assert report["tables"] == 6 and report["pictures"] == 5 and not report["unresolved_tokens"]
    (HERE / "docx_build.json").write_text(
        json.dumps(report, indent=2), encoding="utf-8", newline="\n"
    )
    print(json.dumps(report))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    main(parser.parse_args().output.resolve())
